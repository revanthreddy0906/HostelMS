"""Excel report builders using openpyxl (FR-RG-01/02)."""
from openpyxl import Workbook


def _write_table(ws, headers, rows):
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for col_cells in ws.columns:
        length = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = max(10, length + 2)


def generate_fee_collection_excel(filepath: str, report: dict) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Fee Collection"
    start, end = report["period"]
    ws.append(["Period", f"{start} to {end}"])
    ws.append(["Total Due", report["total_due"]])
    ws.append(["Total Collected", report["total_collected"]])
    ws.append(["Outstanding", report["outstanding"]])
    ws.append([])
    headers = ["Fee ID", "Student ID", "Amount Due", "Amount Paid", "Status", "Due Date"]
    rows = [
        [r["feeid"], r["studentid"], r["amountdue"], r["amountpaid"], r["status"], str(r["duedate"])]
        for r in report["rows"]
    ]
    _write_table(ws, headers, rows)
    wb.save(filepath)
    return filepath


def generate_occupancy_excel(filepath: str, rows: list[dict]) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Occupancy"
    headers = ["Hostel", "Room", "Type", "Capacity", "Occupied", "Free"]
    table_rows = [[r["hostel"], r["room"], r["type"], r["capacity"], r["occupied"], r["free"]] for r in rows]
    _write_table(ws, headers, table_rows)
    wb.save(filepath)
    return filepath


def generate_leave_log_excel(filepath: str, rows: list[dict]) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leave Log"
    headers = ["Leave ID", "Student ID", "Start", "End", "Status", "Exit Logged", "Entry Logged"]
    table_rows = [
        [r["leaveid"], r["studentid"], str(r["startdate"]), str(r["enddate"]), r["status"],
         str(r["exitlogged"] or "-"), str(r["entrylogged"] or "-")]
        for r in rows
    ]
    _write_table(ws, headers, table_rows)
    wb.save(filepath)
    return filepath


def generate_complaint_timeline_excel(filepath: str, rows: list[dict]) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Complaints"
    headers = ["Complaint ID", "Student ID", "Category", "Status", "Created At", "Assigned Staff"]
    table_rows = [
        [r["complaintid"], r["studentid"], r["category"], r["status"], str(r["createdat"]), r["assignedstaffid"] or "-"]
        for r in rows
    ]
    _write_table(ws, headers, table_rows)
    wb.save(filepath)
    return filepath
